"""Phase 2G.2 — JPX Market Data（whole_day xlsx 成交量 / open_interest xlsx 建玉）。

真實檔案結構（2026-09-18 實測）：
- derivatives_market_data_whole_day.xlsx → sheet `market_data_Futures`
  欄位：A=exchange, B=category, C=商品名(日/英), D=session, E=volume, F=J-NET vol, G=value, H=J-NET value
- open_interest.xlsx → sheet `デリバティブ建玉残高状況`（per-contract-month volume + OI）
"""
from __future__ import annotations

import hashlib
import io
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from pydantic import BaseModel

from market_ai_hub.targets.jpx_daily import atomic_write

MARKET_BASE = "https://www.jpx.co.jp/markets/derivatives/trading-volume/tvdivq00000014nn-att"

SESSION_MAP = {"夜間": "night", "前場": "morning", "後場": "afternoon", "合計": "total"}

PRODUCT_EN = {
    "マイクロ": "Nikkei 225 Micro Futures",
    "Nikkei 225 Micro Futures": "Nikkei 225 Micro Futures",
    "Nikkei 225 mini": "Nikkei 225 mini",
    "Nikkei 225 Futures": "Nikkei 225 Futures",
    "日経225先物": "Nikkei 225 Futures",
    "TOPIX": "TOPIX Futures",
}


class ProductVolume(BaseModel):
    product: str
    session: str
    volume: float | None = None
    value: float | None = None
    date: str = ""
    source_url: str = ""


class ContractOI(BaseModel):
    product: str
    contract_month: str
    volume: int | None = None
    open_interest: int | None = None
    oi_change: int | None = None
    date: str = ""
    source_url: str = ""


def _f(v):
    if v is None or v == "" or v == "－":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _i(v):
    f = _f(v)
    return int(f) if f is not None else None


def _product_en(cell: str) -> str | None:
    if not cell:
        return None
    en = cell.split("\n")[-1].strip()
    if "Micro" in en:
        return "Nikkei 225 Micro Futures"
    if "mini" in en:
        return "Nikkei 225 mini"
    if "Nikkei 225 Futures" in en:
        return "Nikkei 225 Futures"
    if "TOPIX" in en:
        return "TOPIX Futures"
    return None


def _session(cell: str) -> str | None:
    if not cell:
        return None
    for jp, en in SESSION_MAP.items():
        if jp in cell:
            return en
    return None


def parse_whole_day_volumes(xlsx_bytes: bytes, date: str = "", source_url: str = "") -> list[ProductVolume]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb["market_data_Futures"]
    out: list[ProductVolume] = []
    current = None
    for row in ws.iter_rows(values_only=True):
        c = row[2] if len(row) > 2 else None
        d = row[3] if len(row) > 3 else None
        e = row[4] if len(row) > 4 else None
        g = row[6] if len(row) > 6 else None
        if c and str(c).strip():
            p = _product_en(str(c))
            if p:
                current = p
        if d and str(d).strip() and current:
            sess = _session(str(d))
            if sess:
                out.append(ProductVolume(product=current, session=sess, volume=_f(e),
                                         value=_f(g), date=date, source_url=source_url))
    return out


def parse_open_interest(xlsx_bytes: bytes, date: str = "", source_url: str = "") -> list[ContractOI]:
    """best-effort 解析 open_interest.xlsx 的 per-contract volume + OI。

    產品 section header 有商品名；契約月 token 為「YYYY年MM月限」。
    只歸屬到日經相關產品（Micro/mini/Large），遇到其他產品 header 即停止歸屬。
    """
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb["デリバティブ建玉残高状況"]
    out: list[ContractOI] = []
    current_product = None
    for row in ws.iter_rows(values_only=True):
        cells = ["" if c is None else str(c).strip() for c in row]
        for cell in cells:
            if not cell:
                continue
            p = _nikkei_product(cell)
            if p:
                current_product = p
                continue
            # 其他產品 header（含 Futures/先物/オプション 但不是日經）→ 停止歸屬
            if _is_non_nikkei_header(cell):
                current_product = None
        if current_product is None:
            continue
        for i, cell in enumerate(cells):
            if cell.endswith("月限"):
                month = _contract_month(cell)
                if month is None:
                    continue
                vol = _i(cells[i + 1]) if i + 1 < len(cells) else None
                oi = _i(cells[i + 2]) if i + 2 < len(cells) else None
                chg = _i(cells[i + 3]) if i + 3 < len(cells) else None
                out.append(ContractOI(product=current_product, contract_month=month,
                                      volume=vol, open_interest=oi, oi_change=chg,
                                      date=date, source_url=source_url))
    return out


def _nikkei_product(cell: str) -> str | None:
    if "日経225マイクロ" in cell or "Nikkei 225 Micro" in cell:
        return "Nikkei 225 Micro Futures"
    if "日経225mini" in cell:
        return "Nikkei 225 mini"
    if "日経225" in cell and "月限" not in cell:
        return "Nikkei 225 Futures"
    return None


def _is_non_nikkei_header(cell: str) -> bool:
    if "月限" in cell:
        return False
    return ("先物" in cell or "Futures" in cell or "オプション" in cell or "Options" in cell)


def _contract_month(cell: str) -> str | None:
    """「2026年12月限」→ 「202612」."""
    import re
    m = re.search(r"(\d{4})年(\d{1,2})月限", cell)
    if not m:
        return None
    return f"{int(m.group(1)):04d}{int(m.group(2)):02d}"


class JPXMarketDataProvider:
    def __init__(self, data_root: Path | None = None) -> None:
        from market_ai_hub.automation.data_lake import DataLakeManager
        self.lake = DataLakeManager(data_root)

    def whole_day_url(self, d: str) -> str:
        return f"{MARKET_BASE}/{d}_derivatives_market_data_whole_day.xlsx"

    def open_interest_url(self, d: str) -> str:
        return f"{MARKET_BASE}/{d}open_interest.xlsx"

    def fetch_whole_day(self, d: str, getter=None) -> list[ProductVolume]:
        if getter is None:
            import httpx
            def getter(u):
                return httpx.get(u, timeout=30, follow_redirects=True).content
        raw = getter(self.whole_day_url(d))
        return parse_whole_day_volumes(raw, date=d, source_url=self.whole_day_url(d))

    def fetch_open_interest(self, d: str, getter=None) -> list[ContractOI]:
        if getter is None:
            import httpx
            def getter(u):
                return httpx.get(u, timeout=30, follow_redirects=True).content
        raw = getter(self.open_interest_url(d))
        return parse_open_interest(raw, date=d, source_url=self.open_interest_url(d))
