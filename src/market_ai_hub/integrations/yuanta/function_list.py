"""Phase 2Y-A — FunctionList.xlsx parser（唯讀，只解析市場 enum + 海外期貨 code）。"""
from __future__ import annotations

from pathlib import Path


def parse_market_enum(xlsx_path: str | Path) -> dict[str, int]:
    """讀 FunctionList 市場類 sheet → {market_name: code}（含 OSE=207）。"""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = None
    for sn in wb.sheetnames:
        cand = wb[sn]
        first = next(cand.iter_rows(values_only=True), None)
        cells = [str(c) for c in (first or []) if c is not None]
        joined = " ".join(cells)
        if "市場" in joined and "代碼" in joined:
            ws = cand
            break
    out: dict[str, int] = {}
    if ws is None:
        return out
    for row in ws.iter_rows(values_only=True):
        vals = ["" if c is None else str(c).strip() for c in row]
        if len(vals) < 2:
            continue
        code = vals[0]
        name = vals[1]
        if code.isdigit():
            out[name] = int(code)
    return out


def find_overseas_futures(xlsx_path: str | Path) -> list[dict]:
    """讀股票代碼總表，回海外期貨（SGX/CME/OSE）列。"""
    return [r for r in load_stock_code_rows(xlsx_path) if r["market_code"] in (202, 203, 207)]


def load_stock_code_rows(xlsx_path: str | Path) -> list[dict]:
    """讀「股票代碼總表」全部列（唯讀）。

    欄位：市場別 | 市場代碼 | 商品代碼(報價 StkCode) | 商品名稱 | 下單代碼 | ...
    """
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = None
    for sn in wb.sheetnames:
        cand = wb[sn]
        first = next(cand.iter_rows(values_only=True), None)
        joined = " ".join(str(c) for c in (first or []) if c is not None)
        if "商品代碼" in joined and "市場" in joined:
            ws = cand
            break
    out: list[dict] = []
    if ws is None:
        return out
    for row in ws.iter_rows(values_only=True):
        vals = ["" if c is None else str(c).strip() for c in row]
        if len(vals) < 4 or not vals[1].isdigit():
            continue
        out.append({
            "market": vals[0], "market_code": int(vals[1]), "code": vals[2],
            "name": vals[3], "order_code": vals[4] if len(vals) > 4 else "",
        })
    return out
